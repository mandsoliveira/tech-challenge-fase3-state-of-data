"""Renderiza o relatório técnico canônico em DOCX e PDF.

A fonte editável fica em ``docs/RELATORIO_TECNICO.md``. O renderizador suporta
os elementos usados pelo relatório: títulos, parágrafos, listas, tabelas,
imagens e o marcador dinâmico ``{{PRINTS_AWS}}``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

RAIZ = Path(__file__).resolve().parents[1]
FONTE = RAIZ / "docs" / "RELATORIO_TECNICO.md"
PASTA_OUTPUT = RAIZ / "output"
PASTA_PRINTS = PASTA_OUTPUT / "prints"
DESTINO_DOCX = PASTA_OUTPUT / "Relatorio_Tecnico_State_of_Data_Fase3.docx"
DESTINO_PDF = PASTA_OUTPUT / "Relatorio_Tecnico_State_of_Data_Fase3.pdf"

AZUL = "17365D"
LARANJA = "E67E22"
GRAFITE = "273746"
CINZA = "6B7280"
CINZA_CLARO = "F3F4F6"


@dataclass(frozen=True)
class Elemento:
    """Elemento intermediário independente do formato de saída."""

    tipo: str
    conteudo: object
    nivel: int = 0


def _limpar_inline(texto: str) -> str:
    """Remove marcação Markdown inline sem alterar o conteúdo."""
    texto = texto.replace("  ", " ").replace("`", "")
    texto = re.sub(r"\*\*(.+?)\*\*", r"\1", texto)
    texto = re.sub(r"\*(.+?)\*", r"\1", texto)
    return texto.strip()


def _separar_celulas(linha: str) -> List[str]:
    return [_limpar_inline(valor.strip()) for valor in linha.strip().strip("|").split("|")]


def _eh_separador_tabela(linha: str) -> bool:
    celulas = _separar_celulas(linha)
    return bool(celulas) and all(re.fullmatch(r":?-{3,}:?", celula) for celula in celulas)


def _imagens_prints() -> List[Path]:
    extensoes = {".png", ".jpg", ".jpeg"}
    return sorted(
        caminho for caminho in PASTA_PRINTS.glob("*.*")
        if caminho.suffix.lower() in extensoes
    )


def analisar_markdown(texto: str) -> List[Elemento]:
    """Converte o subconjunto Markdown do relatório em elementos."""
    linhas = texto.splitlines()
    elementos: List[Elemento] = []
    indice = 0

    while indice < len(linhas):
        linha = linhas[indice].strip()
        if not linha:
            indice += 1
            continue

        if linha == "{{PRINTS_AWS}}":
            prints = _imagens_prints()
            if prints:
                legendas = {
                    "01_s3_camadas_medallion": "Amazon S3 — camadas Bronze, Silver e Gold, além dos resultados do Athena (bucket anonimizado).",
                    "02_glue_jobs_sucesso": "AWS Glue — dois jobs ETL concluídos com sucesso, usando dois workers G.1X cada.",
                    "03_glue_crawler": "AWS Glue Crawler — duas execuções finais concluídas; a tentativa inicial foi corrigida durante a configuração dos alvos.",
                    "04_glue_data_catalog": "Glue Data Catalog — 11 tabelas no banco sod_fase3: 3 Silver e 8 Gold (bucket anonimizado).",
                    "05_athena_resultado_perfil": "Amazon Athena — consulta de perfil de mercado concluída, com percentual, denominador e taxa de resposta.",
                }
                for caminho in prints:
                    legenda = legendas.get(
                        caminho.stem,
                        f"Evidência AWS — {caminho.stem.replace('_', ' ')}",
                    )
                    elementos.append(Elemento("imagem", (caminho, legenda)))
            else:
                elementos.append(Elemento(
                    "paragrafo",
                    "As capturas do console AWS devem ser salvas em output/prints após ocultar o identificador da conta. "
                    "As evidências JSON estruturadas já acompanham a entrega.",
                ))
            indice += 1
            continue

        imagem = re.fullmatch(r"!\[(.+)]\((.+)\)", linha)
        if imagem:
            caminho = (FONTE.parent / imagem.group(2)).resolve()
            elementos.append(Elemento("imagem", (caminho, imagem.group(1))))
            indice += 1
            continue

        titulo = re.match(r"^(#{1,3})\s+(.+)$", linha)
        if titulo:
            elementos.append(Elemento("titulo", _limpar_inline(titulo.group(2)), len(titulo.group(1))))
            indice += 1
            continue

        if linha.startswith("- "):
            itens: List[str] = []
            while indice < len(linhas) and linhas[indice].strip().startswith("- "):
                itens.append(_limpar_inline(linhas[indice].strip()[2:]))
                indice += 1
            elementos.append(Elemento("lista", itens))
            continue

        if linha.startswith("|") and indice + 1 < len(linhas) and _eh_separador_tabela(linhas[indice + 1]):
            cabecalho = _separar_celulas(linha)
            indice += 2
            dados: List[List[str]] = []
            while indice < len(linhas) and linhas[indice].strip().startswith("|"):
                dados.append(_separar_celulas(linhas[indice]))
                indice += 1
            elementos.append(Elemento("tabela", (cabecalho, dados)))
            continue

        partes = [linha]
        indice += 1
        while indice < len(linhas):
            proxima = linhas[indice].strip()
            if not proxima or proxima.startswith(("#", "- ", "|", "![")) or proxima == "{{PRINTS_AWS}}":
                break
            partes.append(proxima)
            indice += 1
        elementos.append(Elemento("paragrafo", _limpar_inline(" ".join(partes))))

    return elementos


def _sombrear(celula, cor: str) -> None:
    propriedades = celula._tc.get_or_add_tcPr()
    sombreado = propriedades.find(qn("w:shd"))
    if sombreado is None:
        sombreado = OxmlElement("w:shd")
        propriedades.append(sombreado)
    sombreado.set(qn("w:fill"), cor)


def _configurar_docx(documento: Document) -> None:
    secao = documento.sections[0]
    secao.top_margin = Cm(1.8)
    secao.bottom_margin = Cm(1.8)
    secao.left_margin = Cm(2.1)
    secao.right_margin = Cm(2.1)

    normal = documento.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(GRAFITE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for nome, tamanho, cor in [
        ("Title", 28, AZUL),
        ("Heading 1", 17, AZUL),
        ("Heading 2", 13, LARANJA),
        ("Heading 3", 11.5, GRAFITE),
    ]:
        estilo = documento.styles[nome]
        estilo.font.name = "Aptos Display"
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string(cor)


def _capa_docx(documento: Document) -> None:
    documento.add_paragraph("\n\n\n")
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("TECH CHALLENGE — FASE 3")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(AZUL)

    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitulo.add_run("Plataforma Analítica State of Data Brasil")
    run.bold = True
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor.from_string(LARANJA)

    documento.add_paragraph("\n\n")
    info = documento.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run("Relatório técnico de Engenharia de Dados e Analytics em AWS\n\n").bold = True
    info.add_run("POSTECH DTAT — Data Analytics\nEquipe do Tech Challenge\nSetembro de 2026")
    documento.add_page_break()


def gerar_docx(elementos: Sequence[Elemento], destino: Path = DESTINO_DOCX) -> Path:
    """Gera a versão editável em DOCX."""
    documento = Document()
    propriedades = documento.core_properties
    propriedades.title = "Relatório Técnico — State of Data Brasil"
    propriedades.subject = "Tech Challenge Fase 3 — Engenharia de Dados e Analytics em AWS"
    propriedades.author = "Equipe do Tech Challenge"
    propriedades.last_modified_by = "Equipe do Tech Challenge"
    propriedades.category = "Relatório acadêmico"
    propriedades.keywords = "State of Data, AWS, Engenharia de Dados, Analytics"
    propriedades.comments = "Gerado de forma reproduzível a partir da fonte em Markdown."
    _configurar_docx(documento)
    _capa_docx(documento)

    for elemento in elementos:
        if elemento.tipo == "titulo":
            if elemento.nivel == 1:  # o título principal já está na capa
                continue
            documento.add_heading(str(elemento.conteudo), level=min(elemento.nivel - 1, 3))
        elif elemento.tipo == "paragrafo":
            paragrafo = documento.add_paragraph(str(elemento.conteudo))
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif elemento.tipo == "lista":
            for item in elemento.conteudo:
                documento.add_paragraph(str(item), style="List Bullet")
        elif elemento.tipo == "tabela":
            cabecalho, dados = elemento.conteudo
            tabela = documento.add_table(rows=1, cols=len(cabecalho))
            tabela.style = "Table Grid"
            for coluna, valor in enumerate(cabecalho):
                celula = tabela.rows[0].cells[coluna]
                celula.text = valor
                _sombrear(celula, AZUL)
                for run in celula.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            for numero_linha, linha in enumerate(dados):
                celulas = tabela.add_row().cells
                for coluna, valor in enumerate(linha):
                    celulas[coluna].text = valor
                    if numero_linha % 2:
                        _sombrear(celulas[coluna], CINZA_CLARO)
            documento.add_paragraph()
        elif elemento.tipo == "imagem":
            caminho, legenda = elemento.conteudo
            if not caminho.exists():
                raise FileNotFoundError(f"Imagem do relatório não encontrada: {caminho}")
            paragrafo = documento.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.add_run().add_picture(str(caminho), width=Cm(16.4))
            texto_legenda = documento.add_paragraph(str(legenda))
            texto_legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in texto_legenda.runs:
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor.from_string(CINZA)

    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(destino)
    return destino


def _estilos_pdf() -> dict:
    base = getSampleStyleSheet()
    return {
        "capa": ParagraphStyle("Capa", parent=base["Title"], fontName="Helvetica-Bold", fontSize=25, leading=30, alignment=TA_CENTER, textColor=colors.HexColor(f"#{AZUL}")),
        "subcapa": ParagraphStyle("Subcapa", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=18, leading=23, alignment=TA_CENTER, textColor=colors.HexColor(f"#{LARANJA}")),
        "h1": ParagraphStyle("H1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor(f"#{AZUL}"), keepWithNext=True),
        "h2": ParagraphStyle("H2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor(f"#{LARANJA}"), keepWithNext=True),
        "corpo": ParagraphStyle("Corpo", parent=base["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.6, alignment=TA_JUSTIFY, textColor=colors.HexColor(f"#{GRAFITE}"), spaceAfter=5),
        "lista": ParagraphStyle("Lista", parent=base["BodyText"], fontName="Helvetica", fontSize=9, leading=12, leftIndent=13, firstLineIndent=-7, textColor=colors.HexColor(f"#{GRAFITE}"), spaceAfter=3),
        "legenda": ParagraphStyle("Legenda", parent=base["BodyText"], fontName="Helvetica-Oblique", fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.HexColor(f"#{CINZA}"), spaceAfter=7),
        "cabecalho": ParagraphStyle("Cabecalho", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=7.8, leading=9.5, textColor=colors.white),
        "celula": ParagraphStyle("Celula", parent=base["BodyText"], fontName="Helvetica", fontSize=7.7, leading=9.5, textColor=colors.HexColor(f"#{GRAFITE}")),
    }


def _rodape_pdf(canvas, documento) -> None:
    canvas.saveState()
    largura, _ = A4
    canvas.setStrokeColor(colors.HexColor(f"#{AZUL}"))
    canvas.line(2 * cm, 1.35 * cm, largura - 2 * cm, 1.35 * cm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor(f"#{CINZA}"))
    canvas.drawString(2 * cm, 0.9 * cm, "Tech Challenge Fase 3 — State of Data Brasil")
    canvas.drawRightString(largura - 2 * cm, 0.9 * cm, f"Página {documento.page}")
    canvas.restoreState()


def _imagem_pdf(caminho: Path, largura_maxima: float = 16.5 * cm, altura_maxima: float = 18.5 * cm) -> Image:
    imagem = Image(str(caminho))
    escala = min(largura_maxima / imagem.imageWidth, altura_maxima / imagem.imageHeight)
    imagem.drawWidth = imagem.imageWidth * escala
    imagem.drawHeight = imagem.imageHeight * escala
    imagem.hAlign = "CENTER"
    return imagem


def gerar_pdf(elementos: Sequence[Elemento], destino: Path = DESTINO_PDF) -> Path:
    """Gera o PDF diretamente, sem depender de LibreOffice ou Word."""
    estilos = _estilos_pdf()
    documento = BaseDocTemplate(
        str(destino),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.7 * cm,
        bottomMargin=1.8 * cm,
        title="Relatório Técnico — State of Data Brasil",
        author="Equipe do Tech Challenge",
        subject="Engenharia de Dados e Analytics em AWS",
        creator="Projeto State of Data Brasil",
        keywords="State of Data, AWS, Engenharia de Dados, Analytics",
    )
    quadro = Frame(documento.leftMargin, documento.bottomMargin, documento.width, documento.height, id="conteudo")
    documento.addPageTemplates([PageTemplate(id="padrao", frames=[quadro], onPage=_rodape_pdf)])

    historia: List[object] = [
        Spacer(1, 4.2 * cm),
        Paragraph("TECH CHALLENGE — FASE 3", estilos["capa"]),
        Spacer(1, 0.6 * cm),
        Paragraph("Plataforma Analítica State of Data Brasil", estilos["subcapa"]),
        Spacer(1, 1.5 * cm),
        Paragraph("Relatório técnico de Engenharia de Dados e Analytics em AWS", estilos["subcapa"]),
        Spacer(1, 2.2 * cm),
        Paragraph("POSTECH DTAT — Data Analytics<br/>Equipe do Tech Challenge<br/>Setembro de 2026", estilos["corpo"]),
        PageBreak(),
    ]

    for elemento in elementos:
        if elemento.tipo == "titulo":
            if elemento.nivel == 1:
                continue
            estilo = estilos["h1"] if elemento.nivel == 2 else estilos["h2"]
            historia.append(Paragraph(str(elemento.conteudo), estilo))
        elif elemento.tipo == "paragrafo":
            historia.append(Paragraph(str(elemento.conteudo), estilos["corpo"]))
        elif elemento.tipo == "lista":
            for item in elemento.conteudo:
                historia.append(Paragraph(f"• {item}", estilos["lista"]))
        elif elemento.tipo == "tabela":
            cabecalho, dados = elemento.conteudo
            linhas = [
                [Paragraph(valor, estilos["cabecalho"]) for valor in cabecalho],
                *[[Paragraph(valor, estilos["celula"]) for valor in linha] for linha in dados],
            ]
            larguras = [documento.width / len(cabecalho)] * len(cabecalho)
            tabela = Table(linhas, colWidths=larguras, repeatRows=1, hAlign="CENTER")
            tabela.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C2CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{CINZA_CLARO}")]),
            ]))
            historia.extend([tabela, Spacer(1, 0.3 * cm)])
        elif elemento.tipo == "imagem":
            caminho, legenda = elemento.conteudo
            if not caminho.exists():
                raise FileNotFoundError(f"Imagem do relatório não encontrada: {caminho}")
            historia.extend([_imagem_pdf(caminho), Paragraph(str(legenda), estilos["legenda"])])

    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.build(historia)
    return destino


def gerar() -> tuple[Path, Path]:
    """Gera os dois formatos e retorna seus caminhos."""
    if not FONTE.exists():
        raise FileNotFoundError(f"Fonte do relatório não encontrada: {FONTE}")
    elementos = analisar_markdown(FONTE.read_text(encoding="utf-8"))
    return gerar_docx(elementos), gerar_pdf(elementos)


def main() -> None:
    docx, pdf = gerar()
    print(f"Relatório DOCX: {docx}")
    print(f"Relatório PDF:  {pdf}")


if __name__ == "__main__":
    main()
